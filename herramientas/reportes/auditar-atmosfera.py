"""Contrasta la exportación real con sus respuestas públicas, sin modificar los archivos auditados."""
import hashlib
import json
import argparse
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('directorio', type=Path)
parser.add_argument('--pdf', type=Path)
parser.add_argument('--excel', type=Path)
args = parser.parse_args()
root = args.directorio
pdf = args.pdf or root / 'atmosfera.pdf'
excel = args.excel or root / 'atmosfera.xlsx'
report = json.loads((root / 'reporte.json').read_text(encoding='utf-8'))
receipt = json.loads((root / 'recibo.json').read_text(encoding='utf-8'))
rows = report['contenido']['filas']
sources = {group: json.loads((root / f'proveedor-{group}.json').read_text(encoding='utf-8')) for group in ['tiempo', 'aire']}
for source in receipt['proveedores']:
    assert hashlib.sha256((root / f"proveedor-{source['grupo']}.json").read_bytes()).hexdigest() == source['sha256']
for row in rows:
    group = 'aire' if row['metrica'].startswith('aqi') else 'tiempo'
    field = {'temperatura': 'temperature_2m', 'aqi': 'us_aqi', 'aqi_europeo': 'european_aqi'}[row['metrica']]
    instant = int(datetime.fromisoformat(row['instante']).timestamp())
    index = sources[group]['hourly']['time'].index(instant)
    assert row['valor'] == sources[group]['hourly'][field][index]
    assert row['unidad'] == sources[group]['hourly_units'][field]
for summary in report['contenido']['resumen']:
    values = [r['valor'] for r in rows if r['punto'] == summary['punto'] and r['metrica'] == summary['metrica'] and r['fuente'] == summary['fuente']]
    assert summary['cantidad'] == len(values) and summary['minimo'] == min(values) and summary['maximo'] == max(values)
    assert abs(summary['promedio'] - sum(values) / len(values)) <= .000051
book = load_workbook(excel, data_only=False)
sheet = book['Lecturas']
assert sheet.max_row == len(rows) + 1
for number, expected in enumerate(rows, 2):
    for col, key in enumerate(report['contenido']['columnas'], 1):
        value = expected.get({'origen': 'origenNombre', 'fuente': 'fuenteNombre'}.get(key, key))
        assert sheet.cell(number, col).value == value
        assert sheet.cell(number, col).data_type not in ('f', 'e')
assert {r['unidad'] for r in rows if r['metrica'].startswith('aqi')} == {'USAQI', 'EAQI'}
reader = PdfReader(pdf)
text = '\n'.join(page.extract_text() or '' for page in reader.pages)
assert all(t in text for t in ['USAQI', 'EAQI', 'escala europea', 'Estados Unidos'])
for i, page in enumerate(reader.pages, 1):
    assert f'GeoSur - {i} / {len(reader.pages)}' in (page.extract_text() or '')
assert hashlib.sha256(pdf.read_bytes()).hexdigest() == receipt['sha256Pdf']
assert hashlib.sha256(excel.read_bytes()).hexdigest() == receipt['sha256Excel']
result = {'paginasPdf': len(reader.pages), 'hojasExcel': book.sheetnames, 'filasExportadas': len(rows),
          'respuestasYExportacionesCoinciden': True, 'resumenesContrastados': len(report['contenido']['resumen']),
          'escalasSeparadas': True, 'formulasInyectadas': 0, 'numeracionComprobada': True}
(root / 'auditoria.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False))
